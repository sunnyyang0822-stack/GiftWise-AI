# 部署版：我们的整页图 PPT + 备注 + P01/P08 透明超链接；同时合成 8 页 PDF（带链接注释）。
import json, sys, os, shutil, hashlib
sys.path.insert(0, "E:/黑客松/.tools/ppt/pylib")
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE
import fitz
D = "E:/黑客松/.tools/ppt/deck/"
M = "E:/黑客松/.tools/giftwise-materials/docs/materials-deployment-20261101/"
URL = "http://117.50.83.57/"
notes = [x["notes"] for x in sorted(json.load(open(M + "REPLACEMENT-NOTES.json", encoding="utf-8")), key=lambda x: x["slide"])]
SHA = "运行 19b6819fe80787717571beb7cc0d17f89587c2d9；证据 6e92a9e076d5989f7ad510bebf3639a9dd2dadf8；入口 " + URL
notes[0] += "\n\n" + SHA
notes[7] += "\n\n" + SHA
json.dump(notes, open(D + "notes-deploy.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
rects = json.load(open(D + "linkrects.json", encoding="utf-8"))
links = {1: rects["s01"], 8: rects["s08"]}
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
k = 12192000 / 1920
blank = prs.slide_layouts[6]
for i, note in enumerate(notes, 1):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(D + f"out/s0{i}.png", 0, 0, prs.slide_width, prs.slide_height)
    if i in links:
        r = links[i]
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(r["left"] * k)), Emu(int(r["top"] * k)), Emu(int(r["width"] * k)), Emu(int(r["height"] * k)))
        sh.fill.background(); sh.line.fill.background(); sh.shadow.inherit = False
        sh.click_action.hyperlink.address = URL
        sh.name = "公网入口链接"
    s.notes_slide.notes_text_frame.text = note
F = M + "final/01-答辩材料/"
bk = M + "work/ppt-native-backup/"
os.makedirs(bk, exist_ok=True)
for n in ("GiftWise-AI-答辩.pptx", "GiftWise-AI-答辩.pdf"):
    if os.path.exists(F + n) and not os.path.exists(bk + n): shutil.copy2(F + n, bk + n)
prs.save(F + "GiftWise-AI-答辩.pptx")
doc = fitz.open()
W, H = 960, 540
for i in range(1, 9):
    pg = doc.new_page(width=W, height=H)
    pg.insert_image(fitz.Rect(0, 0, W, H), filename=D + f"out/s0{i}.png")
    if i in links:
        r = links[i]; s2 = W / 1920
        pg.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(r["left"] * s2, r["top"] * s2, r["right"] * s2, r["bottom"] * s2), "uri": URL})
doc.set_metadata({"title": "GiftWise AI 答辩", "author": "GiftWise AI"})
doc.save(F + "GiftWise-AI-答辩.pdf", deflate=True, garbage=3)
shutil.copy2(F + "GiftWise-AI-答辩.pptx", "E:/黑客松/GiftWise-AI/presentation/GiftWise-AI-答辩(7).pptx")
for n in ("GiftWise-AI-答辩.pptx", "GiftWise-AI-答辩.pdf"):
    print(n, os.path.getsize(F + n), hashlib.sha256(open(F + n, "rb").read()).hexdigest())
